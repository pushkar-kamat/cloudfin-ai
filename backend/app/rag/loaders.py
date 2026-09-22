from __future__ import annotations

import re
from dataclasses import dataclass, asdict
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook

try:
    import xlrd
except ImportError:  # Allows XLSX support even before the optional legacy-XLS package is installed.
    xlrd = None


@dataclass
class DocumentRecord:
    document: str
    text: str
    file_type: str
    section: str | None = None
    page: int | None = None
    content_type: str = "text"
    origin: str = "core"
    source_path: str | None = None

    def to_dict(self) -> dict:
        return asdict(self)


def clean_text(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pretty_name(path: Path) -> str:
    stem = re.sub(r"^[0-9a-fA-F]{8}_", "", path.stem)
    name = stem.replace("_", " ").replace("-", " ").strip()
    return " ".join(word.capitalize() for word in name.split())


def _markdown_records(path: Path, origin: str) -> list[DocumentRecord]:
    text = clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    if not text:
        return []

    document = _pretty_name(path)
    records: list[DocumentRecord] = []
    section = None
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        body = clean_text("\n\n".join(buffer))
        if body:
            records.append(DocumentRecord(
                document=document,
                text=body,
                file_type=path.suffix.lstrip("."),
                section=section,
                origin=origin,
                source_path=str(path),
            ))
        buffer = []

    for line in text.splitlines():
        match = re.match(r"^#{1,4}\s+(.+)$", line.strip())
        if match:
            flush()
            section = match.group(1).strip()
        else:
            buffer.append(line)
    flush()

    if not records:
        records.append(DocumentRecord(
            document=document,
            text=text,
            file_type=path.suffix.lstrip("."),
            origin=origin,
            source_path=str(path),
        ))
    return records


def _txt_records(path: Path, origin: str) -> list[DocumentRecord]:
    text = clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    if not text:
        return []
    return [DocumentRecord(
        document=_pretty_name(path),
        text=text,
        file_type="txt",
        origin=origin,
        source_path=str(path),
    )]


def _pdf_records(path: Path, origin: str) -> list[DocumentRecord]:
    records: list[DocumentRecord] = []
    document_name = _pretty_name(path)
    pdf = fitz.open(path)
    if pdf.needs_pass:
        pdf.close()
        raise ValueError("Password-protected PDFs are not supported")

    for index, page in enumerate(pdf):
        text = clean_text(page.get_text("text"))
        if not text:
            continue
        first_line = next((x.strip() for x in text.splitlines() if x.strip()), "")
        section = first_line if 3 <= len(first_line) <= 90 else None
        records.append(DocumentRecord(
            document=document_name,
            text=text,
            file_type="pdf",
            section=section,
            page=index + 1,
            origin=origin,
            source_path=str(path),
        ))
    pdf.close()
    return records


def _docx_records(path: Path, origin: str) -> list[DocumentRecord]:
    doc = Document(path)
    document_name = _pretty_name(path)
    records: list[DocumentRecord] = []
    section = None
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        text = clean_text("\n\n".join(buffer))
        if text:
            records.append(DocumentRecord(
                document=document_name,
                text=text,
                file_type="docx",
                section=section,
                origin=origin,
                source_path=str(path),
            ))
        buffer = []

    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style_name.startswith("heading"):
            flush()
            section = text
        else:
            buffer.append(text)
    flush()

    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            values = [clean_text(cell.text) for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
        if rows:
            records.append(DocumentRecord(
                document=document_name,
                text="\n".join(rows),
                file_type="docx",
                section=f"Table {table_index}",
                content_type="table",
                origin=origin,
                source_path=str(path),
            ))
    return records



def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(str(value))


def _rows_to_excel_records(
    path: Path,
    origin: str,
    sheet_name: str,
    rows: list[list[str]],
) -> list[DocumentRecord]:
    """Turn one worksheet into small searchable table records.

    The first non-empty row is treated as the header and repeated in each block so
    retrieved chunks keep their column meaning. This stays lightweight and avoids
    adding pandas just for spreadsheet ingestion.
    """
    non_empty = [row for row in rows if any(cell for cell in row)]
    if not non_empty:
        return []

    document_name = _pretty_name(path)
    header = non_empty[0]
    data_rows = non_empty[1:]
    header_line = " | ".join(cell or f"Column {idx + 1}" for idx, cell in enumerate(header))

    records: list[DocumentRecord] = []
    # Keep row groups modest; the normal chunker can split further if necessary.
    block_size = 20
    blocks = [data_rows[i:i + block_size] for i in range(0, len(data_rows), block_size)] or [[]]

    for block_index, block in enumerate(blocks, start=1):
        lines = [header_line]
        for row in block:
            width = max(len(header), len(row))
            padded = row + [""] * (width - len(row))
            lines.append(" | ".join(padded))

        text = clean_text("\n\n".join(line for line in lines if line.strip(" |")))
        if not text:
            continue
        section = f"Sheet: {sheet_name}"
        if len(blocks) > 1:
            section += f" · Rows {((block_index - 1) * block_size) + 2}-{((block_index - 1) * block_size) + len(block) + 1}"
        records.append(DocumentRecord(
            document=document_name,
            text=text,
            file_type=path.suffix.lstrip("."),
            section=section,
            content_type="table",
            origin=origin,
            source_path=str(path),
        ))
    return records


def _xlsx_records(path: Path, origin: str) -> list[DocumentRecord]:
    records: list[DocumentRecord] = []
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = [
                [_cell_text(value) for value in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            records.extend(_rows_to_excel_records(path, origin, worksheet.title, rows))
    finally:
        workbook.close()
    return records


def _xls_records(path: Path, origin: str) -> list[DocumentRecord]:
    if xlrd is None:
        raise ValueError("Legacy .xls support requires the xlrd package. Run pip install -r requirements.txt")
    records: list[DocumentRecord] = []
    workbook = xlrd.open_workbook(path)
    for sheet in workbook.sheets():
        rows = [
            [_cell_text(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
            for row_idx in range(sheet.nrows)
        ]
        records.extend(_rows_to_excel_records(path, origin, sheet.name, rows))
    return records

def load_document(path: Path, origin: str = "runtime") -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        records = _pdf_records(path, origin)
    elif suffix == ".docx":
        records = _docx_records(path, origin)
    elif suffix == ".md":
        records = _markdown_records(path, origin)
    elif suffix == ".txt":
        records = _txt_records(path, origin)
    elif suffix == ".xlsx":
        records = _xlsx_records(path, origin)
    elif suffix == ".xls":
        records = _xls_records(path, origin)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    return [record.to_dict() for record in records]
